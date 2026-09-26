"""Read the source wordbook and optionally enrich cards from dictionary pages.

Importing this module has no file or network side effects. Network lookups and
audio downloads run only when their respective functions are called.
"""

from __future__ import annotations

import hashlib
import re
import tempfile
import unicodedata
from pathlib import Path
from urllib.parse import urlparse

import requests
from lxml import etree, html
from openpyxl import load_workbook


class InputError(ValueError):
    """The source data is ambiguous or a dictionary response is unusable."""


_MAX_AUDIO_BYTES = 5 * 1024 * 1024
_MIN_AUDIO_BYTES = 128
_AUDIO_HOST = "www.oxfordlearnersdictionaries.com"
_USER_AGENT = "AnkiWordbook/1.0 (+local dictionary enrichment)"


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _position(value: object, context: str) -> str:
    if isinstance(value, bool):
        raise InputError(f"{context}: invalid position {value!r}")
    if isinstance(value, float):
        if not value.is_integer():
            raise InputError(f"{context}: non-integral position {value!r}")
        return str(int(value))
    return _text(value)


def _lookup_word(word: str) -> str:
    first = _text(word).replace("\r\n", "\n").replace("\r", "\n").split("\n", 1)[0].strip()
    if not first:
        raise InputError("word has no first-line lookup term")
    return first


def audio_filename(word: str) -> str:
    """Return a stable, safe MP3 filename based on the word's first line.

    Plain ASCII headwords retain their legacy media names. Other headwords
    carry a short digest to prevent different punctuation from colliding.
    """
    headword = unicodedata.normalize("NFKC", _lookup_word(word))
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", headword).strip("_-")[:80]
    if not safe:
        safe = "word"
    if safe != headword:
        safe += "_" + hashlib.sha256(headword.encode("utf-8")).hexdigest()[:12]
    return safe + ".mp3"


def read_wordbook(path: Path) -> list[dict]:
    """Read a Chinese Excel wordbook without changing it.

    A definition-only row continues the preceding word in the same lesson.
    ``position`` is returned as source text; it is never used to infer a
    translation or a card identity.
    """
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    cards: list[dict] = []
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            if header_row is None:
                continue
            headings = [_text(value) for value in header_row]
            if len(set(name for name in headings if name)) != len([name for name in headings if name]):
                raise InputError(f"{sheet.title}: duplicate column header")
            required = ("Lesson", "序号", "单词", "词义")
            missing = [name for name in required if name not in headings]
            if missing:
                raise InputError(f"{sheet.title}: missing columns {', '.join(missing)}")
            columns = {name: index for index, name in enumerate(headings) if name}

            def cell(row: tuple, *names: str) -> str:
                for name in names:
                    index = columns.get(name)
                    if index is not None and index < len(row):
                        value = _text(row[index])
                        if value:
                            return value
                return ""

            seen_words: set[tuple[str, str]] = set()
            seen_positions: set[tuple[str, str]] = set()
            previous: dict | None = None
            for row_number, row in enumerate(rows, 2):
                if not any(_text(value) for value in row):
                    continue
                context = f"{sheet.title}!{row_number}"
                lesson = cell(row, "Lesson")
                position_index = columns["序号"]
                position = _position(row[position_index] if position_index < len(row) else None, context)
                word = cell(row, "单词")
                definition = cell(row, "词义")
                if word:
                    if not lesson or not position:
                        raise InputError(f"{context}: word requires Lesson and 序号")
                    word_key = (lesson, word.casefold())
                    position_key = (lesson, position)
                    if word_key in seen_words:
                        raise InputError(f"{context}: duplicate word in {lesson}: {word!r}")
                    if position_key in seen_positions:
                        raise InputError(f"{context}: duplicate position in {lesson}: {position!r}")
                    seen_words.add(word_key)
                    seen_positions.add(position_key)
                    previous = {
                        "sheet": sheet.title,
                        "lesson": lesson,
                        "position": position,
                        "word": word,
                        "phonetic": cell(row, "音标", "phonetic"),
                        "definition": definition,
                        "simple_definition": cell(row, "简明释义", "simple_definition"),
                        "level": cell(row, "考试等级", "level"),
                        "word_forms": cell(row, "词形变化", "word_forms"),
                        "audio_filename": audio_filename(word),
                    }
                    cards.append(previous)
                elif definition:
                    if previous is None or position or (lesson and lesson != previous["lesson"]):
                        raise InputError(f"{context}: orphan or ambiguous definition continuation")
                    previous["definition"] += ("\n" if previous["definition"] else "") + definition
                else:
                    raise InputError(f"{context}: nonempty row has no word or definition")
    finally:
        workbook.close()
    return cards


def _document(response: requests.Response, source: str):
    response.raise_for_status()
    try:
        return html.fromstring(response.content)
    except (ValueError, TypeError, etree.LxmlError) as exc:
        raise InputError(f"{source}: response is not parseable HTML") from exc


def _class(name: str) -> str:
    return f"contains(concat(' ', normalize-space(@class), ' '), ' {name} ')"


def _node_text(node) -> str:
    return " ".join(" ".join(node.itertext()).split())


def lookup_oxford(word: str, timeout: float = 10) -> dict:
    """Read US phonetic and MP3 URL from an exact Oxford entry page.

    Empty optional fields are valid on an entry page. A search or error page
    raises ``InputError`` instead of suggesting a potentially different word.
    """
    term = _lookup_word(word)
    response = requests.get(
        "https://www.oxfordlearnersdictionaries.com/search/english/direct/",
        params={"q": term},
        headers={"User-Agent": _USER_AGENT},
        timeout=timeout,
    )
    tree = _document(response, "Oxford")
    final_url = getattr(response, "url", "")
    parsed = urlparse(final_url)
    if parsed.scheme != "https" or parsed.hostname != _AUDIO_HOST or not parsed.path.startswith("/definition/english/"):
        raise InputError(f"Oxford: no exact entry page for {term!r}")
    if not tree.xpath(f"//*[{_class('entry')}] | //*[@id='entryContent']"):
        raise InputError(f"Oxford: unsupported entry markup for {term!r}")
    # The direct-search endpoint can redirect a misspelling to a suggested
    # entry. Do not silently enrich the original word with another headword.
    headwords = tree.xpath(f"//h1[{_class('headword')}]")
    if headwords:
        exact = _node_text(headwords[0]).casefold() == term.casefold()
    else:
        slug = parsed.path.rsplit("/", 1)[-1]
        slug = re.sub(r"_\d+(?:_\d+)?$", "", slug)
        exact = slug.casefold() == re.sub(r"\s+", "-", term.casefold())
    if not exact:
        raise InputError(f"Oxford: redirected to a different headword for {term!r}")
    phonetics = tree.xpath(f"//*[{_class('phons_n_am')}]//*[{_class('phon')}]")
    buttons = tree.xpath(f"//*[{_class('audio_play_button')} and {_class('pron-us')}]")
    phonetic = _node_text(phonetics[0]) if phonetics else ""
    audio_url = _text(buttons[0].get("data-src-mp3")) if buttons else ""
    if audio_url:
        _validate_audio_url(audio_url)
    return {"phonetic": phonetic, "audio_url": audio_url}


def lookup_youdao(word: str, timeout: float = 10) -> dict:
    """Parse static Youdao result markup; report pages requiring JavaScript."""
    term = _lookup_word(word)
    response = requests.get(
        "https://dict.youdao.com/result",
        params={"word": term, "lang": "en"},
        headers={"User-Agent": _USER_AGENT},
        timeout=timeout,
    )
    tree = _document(response, "Youdao")
    definitions = []
    for item in tree.xpath(f"//li[{_class('word-exp')}]"):
        pos = item.xpath(f".//*[{_class('pos')}]")
        trans = item.xpath(f".//*[{_class('trans')}]")
        translation = _node_text(trans[0]) if trans else ""
        if translation:
            definitions.append(" ".join(part for part in (_node_text(pos[0]) if pos else "", translation) if part))
    if not definitions:
        raise InputError(f"Youdao: no static definitions for {term!r}; dynamic page is unsupported")
    levels = [_node_text(node) for node in tree.xpath(f"//*[{_class('exam_type-value')}]")]
    forms = []
    for node in tree.xpath(f"//*[{_class('word-wfs-cell-less')}]"):
        names = node.xpath(f".//*[{_class('wfs-name')}]")
        values = node.xpath(f".//*[{_class('transformation')}]")
        if names and values:
            name, value = _node_text(names[0]), _node_text(values[0])
            if name and value:
                forms.append(f"{name}：{value}")
    return {
        "simple_definition": " | ".join(definitions),
        "level": " | ".join(level for level in levels if level),
        "word_forms": " | ".join(forms),
    }


def _validate_audio_url(url: str) -> None:
    parsed = urlparse(url)
    try:
        port = parsed.port
    except ValueError as exc:
        raise InputError("invalid audio URL port") from exc
    if (
        parsed.scheme != "https"
        or parsed.hostname != _AUDIO_HOST
        or port not in (None, 443)
        or not parsed.path.startswith("/media/english/us_pron/")
        or not parsed.path.endswith(".mp3")
        or parsed.username is not None
        or parsed.password is not None
    ):
        raise InputError("audio URL is not an Oxford US pronunciation MP3")


def _valid_mp3(data: bytes) -> bool:
    if len(data) < _MIN_AUDIO_BYTES:
        return False
    offset = 0
    if data.startswith(b"ID3"):
        if len(data) < 10 or any(byte & 0x80 for byte in data[6:10]):
            return False
        tag_size = sum(byte << shift for byte, shift in zip(data[6:10], (21, 14, 7, 0)))
        offset = 10 + tag_size
    if len(data) < offset + 4:
        return False
    header = int.from_bytes(data[offset:offset + 4], "big")
    return (
        header >> 21 == 0x7FF
        and (header >> 19) & 0b11 != 0b01  # reserved MPEG version
        and (header >> 17) & 0b11 != 0b00  # reserved layer
        and (header >> 12) & 0b1111 not in (0, 15)  # invalid bitrate
        and (header >> 10) & 0b11 != 0b11  # invalid sample rate
    )


def download_audio(url: str, word: str, directory: Path, timeout: float = 10) -> str:
    """Download one bounded MP3 to *directory* atomically; return its filename."""
    _validate_audio_url(url)
    filename = audio_filename(word)
    directory = Path(directory)
    target = directory / filename
    if target.is_symlink():
        raise InputError("audio target must not be a symlink")
    if target.is_file():
        data = target.read_bytes() if target.stat().st_size <= _MAX_AUDIO_BYTES else b""
        if _valid_mp3(data):
            return filename

    response = requests.get(
        url,
        headers={"User-Agent": _USER_AGENT},
        timeout=timeout,
        stream=True,
        allow_redirects=False,
    )
    try:
        response.raise_for_status()
        if 300 <= response.status_code < 400:
            raise InputError("audio download redirects are unsupported")
        content_type = response.headers.get("Content-Type", "").split(";", 1)[0].strip().lower()
        if content_type not in {"audio/mpeg", "audio/mp3", "application/octet-stream"}:
            raise InputError(f"unexpected audio content type: {content_type!r}")
        content_length = response.headers.get("Content-Length")
        if content_length:
            try:
                length = int(content_length)
            except ValueError as exc:
                raise InputError("invalid audio Content-Length") from exc
            if length < 0:
                raise InputError("invalid audio Content-Length")
            if length > _MAX_AUDIO_BYTES:
                raise InputError("audio exceeds size limit")
        data = bytearray()
        for chunk in response.iter_content(chunk_size=64 * 1024):
            if len(chunk) > _MAX_AUDIO_BYTES - len(data):
                raise InputError("audio exceeds size limit")
            data.extend(chunk)
        if not _valid_mp3(data):
            raise InputError("response is not a valid MP3")
        directory.mkdir(parents=True, exist_ok=True)
        temp_name = None
        try:
            with tempfile.NamedTemporaryFile(mode="wb", prefix=".audio-", dir=directory, delete=False) as temp:
                temp_name = temp.name
                temp.write(data)
            Path(temp_name).replace(target)
        finally:
            if temp_name is not None:
                Path(temp_name).unlink(missing_ok=True)
        return filename
    finally:
        response.close()
